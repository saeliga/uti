#!/usr/bin/env python3
"""
Gera js/impressos/modelos.js a partir das planilhas da Liga (somente o LEIAUTE:
larguras, alturas, mesclagens, bordas, fontes e rótulos). Nenhum dado de paciente
é copiado: células com fórmula viram marcadores {{campo}} e as abas usadas são a
GRADE (ocorrência em branco) e a "BH LEITO 01" (rótulos fixos + fórmulas).

Uso: python3 gerar_impressos.py <Ocorrencia_e_Balanco_Hidrico.xlsx> <saida.js>
"""
import datetime
import html
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

ORIGEM, SAIDA = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(ORIGEM)

# ---------------------------------------------------------------------------
# Marcadores: célula → texto (com {{campo}}). Substituem fórmulas e campos de dado.
MARC_BH = {
    'A1': 'FOLHA DE CUIDADOS - UTI GERAL', 'AV1': '{{dataExtenso}}',
    'U2': 'PACIENTE: {{nome}}', 'BL2': '{{leito}}',
    'U3': 'REG: {{registro}}', 'AF3': 'CONVÊNIO: {{convenio}}', 'AS3': 'DATA NASCIMENTO: {{nascimento}}',
    'U4': 'TIPO ISOLAMENTO: {{isolamento}}', 'AQ4': 'ALERGIAS: {{alergias}}',
    'A5': 'DIAGNÓSTICO: {{diagnostico}}', 'A6': '{{diagnosticoCirurgico}}',
    'AV42': '{{dataExtenso}}', 'U43': 'PACIENTE: {{nome}}', 'BL43': '{{leito}}',
    'U44': 'REG: {{registro}}', 'AF44': 'CONVÊNIO: {{convenio}}', 'AS44': 'DATA NASCIMENTO: {{nascimento}}',
    'U45': 'TIPO ISOLAMENTO: {{isolamento}}', 'AQ45': 'ALERGIAS: {{alergias}}',
    'A46': 'DIAGNÓSTICO: {{diagnostico}}', 'A47': '{{diagnosticoCirurgico}}',
    # Procedimentos/invasões e dias de permanência
    'AD50': '{{diasUti}}',
    'W51': 'AVP', 'AD51': '{{diasAvp}}',
    'W52': 'AVC', 'AD52': '{{diasAvc}}',
    'W53': 'CAT HD', 'AD53': '{{diasCatHd}}',
    'W54': '{{sondaRotulo}}', 'AD54': '{{diasSonda}}',
    'AL80': '0%', 'AT80': '50%', 'BB80': '75%', 'BI80': '100%',
    # Aceitação da dieta (prato e garrafa), para o técnico marcar à mão
    'AH79': '{{prato0}}', 'AJ79': '{{garrafa0}}', 'AP79': '{{prato50}}', 'AR79': '{{garrafa50}}',
    'AX79': '{{prato75}}', 'AZ79': '{{garrafa75}}', 'BE79': '{{prato100}}', 'BG79': '{{garrafa100}}',
    'AG50': '', 'AM50': '', 'A2': '{{logo}}', 'A43': '{{logo}}',
    'AM51': '{{diasVmi}}', 'AM52': '{{diasTqt}}', 'AM53': '{{diasTuboLaminar}}', 'AM54': '{{diasDreno}}',
    'AG55': '{{svdRotulo}}', 'AM55': '{{diasSvd}}',
    'AD55': '{{diasGtm}}', 'AD56': '{{diasJejuno}}', 'AM56': '{{diasUrostomia}}',
    # Antimicrobianos (até 4) e dias
    **{f'AP{49 + k}': f'{{{{atb{k}}}}}' for k in range(4)},
    **{f'BJ{49 + k}': f'{{{{atbDias{k}}}}}' for k in range(4)},
}
MARC_OC = {
    'A1': 'OCORRÊNCIA  DE ENFERMAGEM  - UTI - POL',
    'A3': 'TIPO ISOLAMENTO: {{isolamento}}', 'H3': '{{data}}',
    'B4': 'LEITO {{leitoNum}}', 'D4': '{{xClinico}}', 'F4': '{{xCirurgico}}',
    'A5': 'PACIENTE:  {{nome}}', 'G5': 'IDADE: {{idade}}',
    'A6': 'DATA NASCIMENTO: {{nascimento}}', 'C6': 'PULSEIRA: ({{pulseiraS}}) S ({{pulseiraN}}) N         CX. ESCOVA (   )', 'G6': 'REG: {{registro}}',
    'A7': 'CONVÊNIO: {{convenio}}', 'B7': 'DIH: {{dih}}', 'G7': '{{duti}}', 'H7': '{{diasUti}}',
    'A8': 'DIAGNÓSTICO: {{diagnostico}}', 'A9': 'DIAGNÓSTICO CIRÚRGICO: ',
    'A10': 'MÉDICOS RESPONSÁVEIS: {{medicos}}', 'E10': 'ALERGIAS: {{alergias}}',
    'E12': '{{boneco}}<div class="oc-lesoes">{{lesoes}}</div>',
    'F27': 'DATA {{sondaRotulo}}:', 'H27': '{{dataSonda}}',
    'E28': '({{viaVo}}) VO ({{viaSne}}) SNE ({{viaSoe}}) SOE ({{viaSng}}) SNG ({{viaGtm}}) GTM (   ) JTM ({{viaNpt}}) NPT',
    'B34': '{{avpData}}', 'C34': '{{avpLocal}}', 'B35': '{{avcData}}', 'C35': '{{avcLocal}}',
    'B36': '{{ctiData}}', 'C36': '{{ctiLocal}}', 'B37': '{{catHdData}}', 'C37': '{{catHdLocal}}',
    'G36': 'DATA {{svdRotulo}}:', 'H36': '{{dataSvd}}',
    'E37': '({{eEsp}}) ESPONT ({{eSvd}}) SVD ({{eSva}}) SVA (   ) DISP. URINÁRIO ({{eFralda}}) FRALDA',
    'A39': 'HV {{vHV}}', 'B39': 'NORADRENALINA {{vNORADRENALINA}}',
    'A40': 'MIDAZOLAM {{vMIDAZOLAM}}', 'B40': 'DOBUTAMINA {{vDOBUTAMINA}}',
    'A41': 'FENTANIL {{vFENTANIL}}', 'B41': 'DOPAMINA {{vDOPAMINA}}',
    'A42': 'SG 10% {{vSG10}}', 'B42': 'NITROPRUSSIATO {{vNITROPRUSSIATO}}',
    'A43': 'INSULINA {{vINSULINA}}', 'B43': 'NITROGLICERINA {{vNITROGLICERINA}}',
    'A44': 'AMIODARONA {{vAMIODARONA}}', 'B44': '{{outraInf0}}', 'A45': '{{outraInf1}}', 'B45': '{{outraInf2}}',
    **{f'A{47 + k}': f'{{{{atb{k}}}}}' for k in range(5)},
    **{f'B{47 + k}': f'{{{{atbDias{k}}}}}' for k in range(5)},
    'E42': '(    ) ESPONT ({{eColost}}) COLOSTOMIA ({{eIleost}}) ILEOSTOMIA (   ) FRALDA',
    'E50': '({{dPenrose}}) PENROSE ({{dLaminar}}) TUBO LAMINAR ({{dToracico}}) TORÁCICO ({{dOutro}}) {{dOutroNome}}',
    'A53': '({{cAvc}}) AVC {{dCurAvc}} (   ) DISS VENOSA ({{cDreno}}) DRENO',
    'A54': '({{cCvti}}) CVTI {{dCurCvti}} ({{cCirurgico}})  CIRÚRGICO   ({{cLpp}})  LESÃO P/ PRESSÃO',
    # Abaixo de BIPAP: via aérea artificial (TOT ou TQT) e data de inserção
    'A27': '{{viaAerea}}',
    # Ocorrências: culturas e ocorrências abertas da evolução, uma por linha
    'E51': 'OCORRÊNCIAS: {{oc0}}', **{f'E{51 + k}': f'{{{{oc{k}}}}}' for k in range(1, 6)},
}
# Mesclagens que faltam na planilha (no Excel o texto transborda; na tabela, não).
MESCLAR_OC = ['E40:H40']

# ---------------------------------------------------------------------------
ESTILOS = {}  # estilo inline → classe (deixa o modelo bem menor)


def classe(st, prefixo):
    if not st:
        return ''
    if st not in ESTILOS:
        ESTILOS[st] = f'{prefixo}{len(ESTILOS)}'
    return ESTILOS[st]


BORDA = {'thin': '0.5pt solid #000', 'hair': '0.5pt solid #000', 'medium': '1.2pt solid #000', 'thick': '1.8pt solid #000',
         'dotted': '0.5pt dotted #000', 'dashed': '0.5pt dashed #000', 'double': '2pt double #000'}


def texto_celula(c):
    v = c.value
    if v is None or isinstance(v, bool):
        return ''
    if isinstance(v, datetime.datetime):
        return f'{v.hour}:{v.minute:02d}'
    if isinstance(v, datetime.time):
        return f'{v.hour}:{v.minute:02d}'
    if isinstance(v, float):
        if c.number_format and '%' in c.number_format:
            return f'{round(v * 100)}%'
        return str(v)
    s = str(v)
    if s.startswith('='):
        return ''  # fórmula sem marcador: fica vazia (nunca copia dado)
    return s


def estilo(c, linha_altura_pt):
    s = []
    b = c.border
    for lado, nome in (('top', 'top'), ('right', 'right'), ('bottom', 'bottom'), ('left', 'left')):
        st = getattr(getattr(b, lado), 'style', None)
        if st:
            s.append(f'border-{nome}:{BORDA.get(st, BORDA["thin"])}')
    f = c.font
    if f is not None:
        if f.b or (f.name or '').lower().endswith('black'):
            s.append('font-weight:700')
        if f.i:
            s.append('font-style:italic')
        if f.sz:
            s.append(f'font-size:{float(f.sz):g}pt')
    if c.fill is not None and c.fill.fill_type == 'solid':
        rgb = c.fill.fgColor.rgb if isinstance(c.fill.fgColor.rgb, str) else None
        if rgb and rgb not in ('00000000', 'FFFFFFFF'):
            s.append(f'background:#{rgb[-6:]}')
            if rgb[-6:].upper() == '000000':
                s.append('color:#fff')
    a = c.alignment
    if a is not None:
        if a.horizontal in ('center', 'centerContinuous'):
            s.append('text-align:center')
        elif a.horizontal == 'right':
            s.append('text-align:right')
        if a.vertical == 'top':
            s.append('vertical-align:top')
        elif a.vertical == 'center':
            s.append('vertical-align:middle')
        if a.textRotation in (90, 180, 255):
            s.append('writing-mode:vertical-rl;transform:rotate(180deg)')
        if a.wrap_text:
            s.append('white-space:normal')
    return ';'.join(s)


def tabela(ws, linhas, colunas, marc, larguras, altura_pt, regiao_sem_borda=()):
    ncol = len(colunas)
    mapa = {}
    for m in ws.merged_cells.ranges:
        mapa[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        for r in range(m.min_row, m.max_row + 1):
            for cc in range(m.min_col, m.max_col + 1):
                if (r, cc) != (m.min_row, m.min_col):
                    mapa[(r, cc)] = None
    total = sum(larguras)
    partes = ['<table class="planilha"><colgroup>']
    partes += [f'<col style="width:{w / total * 100:.3f}%">' for w in larguras]
    partes.append('</colgroup>')
    for r in linhas:
        partes.append(f'<tr style="height:{altura_pt(r):.2f}pt">')
        for k, cc in enumerate(colunas):
            if (r, cc) in mapa and mapa[(r, cc)] is None:
                continue
            span = mapa.get((r, cc))
            rs, cs = (span or (1, 1))
            # mesclagens que ultrapassam a faixa de linhas desta página
            rs = min(rs, linhas[-1] - r + 1)
            cs = min(cs, ncol - k)
            c = ws.cell(r, cc)
            coord = f'{get_column_letter(cc)}{r}'
            txt = marc[coord] if coord in marc else html.escape(texto_celula(c))
            st = estilo(c, altura_pt(r))
            # A borda da direita/baixo de uma mesclagem vem da última célula
            if span:
                fim = ws.cell(r + rs - 1, cc + cs - 1)
                for lado in ('right', 'bottom'):
                    stl = getattr(getattr(fim.border, lado), 'style', None)
                    if stl and f'border-{lado}' not in st:
                        st += f';border-{lado}:{BORDA.get(stl, BORDA["thin"])}'
            attrs = ''
            if rs > 1:
                attrs += f' rowspan="{rs}"'
            if cs > 1:
                attrs += f' colspan="{cs}"'
            if st:
                attrs += f' class="{classe(st, "e")}"'
            if coord in marc:
                attrs += f' data-c="{coord}"'
            partes.append(f'<td{attrs}>{txt}</td>')
        partes.append('</tr>')
    partes.append('</table>')
    return ''.join(partes)


def largura_col(ws, letra, padrao=12.63):
    idx = column_index_from_string(letra)
    for k, v in ws.column_dimensions.items():
        if v.min and v.max and v.min <= idx <= v.max and v.width:
            return v.width
    return padrao


# BH: 66 colunas iguais; página 1 = linhas 1 a 41; página 2 = 42 a 80.
bh = wb['BH LEITO 01']
cols_bh = list(range(1, 67))
larg_bh = [largura_col(bh, get_column_letter(c)) for c in cols_bh]
alt_bh = lambda r: bh.row_dimensions[r].height or 15.75
bh1 = tabela(bh, list(range(1, 42)), cols_bh, MARC_BH, larg_bh, alt_bh)
bh2 = tabela(bh, list(range(42, 81)), cols_bh, MARC_BH, larg_bh, alt_bh)

# Ocorrência: GRADE, colunas A a H, linhas 1 a 56.
oc = wb['GRADE']
for faixa in MESCLAR_OC:
    oc.merge_cells(faixa)
cols_oc = list(range(1, 9))
larg_oc = [largura_col(oc, get_column_letter(c)) for c in cols_oc]
alt_oc = lambda r: oc.row_dimensions[r].height or 15.75
oco = tabela(oc, list(range(1, 57)), cols_oc, MARC_OC, larg_oc, alt_oc)

js = (
    '// GERADO por scripts/gerar_impressos.py a partir das planilhas da Liga (somente o\n'
    '// leiaute e os rótulos; nenhum dado de paciente). Não edite à mão: rode o gerador.\n'
    '// Marcadores {{campo}} são preenchidos em js/impressos/*.js.\n\n'
    f'export const BH_PAGINA_1 = {json.dumps(bh1, ensure_ascii=False)};\n\n'
    f'export const BH_PAGINA_2 = {json.dumps(bh2, ensure_ascii=False)};\n\n'
    f'export const OCORRENCIA = {json.dumps(oco, ensure_ascii=False)};\n\n'
    f'export const ESTILOS_PLANILHA = {json.dumps("".join(f".planilha .{c}{{{st}}}" for st, c in ESTILOS.items()), ensure_ascii=False)};\n'
)
open(SAIDA, 'w', encoding='utf-8').write(js)
print('ok', len(bh1), len(bh2), len(oco))
